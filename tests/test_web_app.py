"""Web app tests: upload, background extraction, job/deck status endpoints,
style guide reading and editing, outline planning and restart recovery.
Everything runs against temporary directories, so no repository artifacts
are touched."""

from __future__ import annotations

import json
import time
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import config, db
from app.generation import run_generation
from app.main import create_app
from core.critique import CritiqueReport
from core.generation import DeckBrief, GeneratedDeck, GenerationPlan
from core.style_guide import StyleGuide
from core.templates import load_templates

POLL_TIMEOUT = 60.0


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "UPLOADS_DIR", tmp_path / "uploads")
    monkeypatch.setattr(config, "GENERATED_DIR", tmp_path / "generated")
    monkeypatch.setattr(config, "STYLE_GUIDES_DIR", tmp_path / "style_guides")
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "deckdna.sqlite3")
    with TestClient(create_app()) as test_client:
        yield test_client


def _upload_deck(client, fixture_deck, name="sample_deck.pptx"):
    response = client.post(
        "/v1/decks",
        files={
            "file": (
                name,
                fixture_deck.read_bytes(),
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation",
            )
        },
    )
    return response


def _wait_for_job(client, job_id, timeout=POLL_TIMEOUT):
    deadline = time.monotonic() + timeout
    while True:
        response = client.get(f"/v1/jobs/{job_id}")
        assert response.status_code == 200, response.text
        job = response.json()
        if job["status"] in ("completed", "failed"):
            return job
        assert time.monotonic() < deadline, f"job did not finish in time: {job}"
        time.sleep(0.05)


def test_upload_runs_extraction_and_reports_deck(client, fixture_deck):
    response = _upload_deck(client, fixture_deck)
    assert response.status_code == 202, response.text
    created = response.json()
    assert created["status"] == "queued"
    assert created["deck_id"].startswith("deck_")
    assert created["job_id"].startswith("job_")

    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job
    assert job["progress"] == 100
    assert job["error"] is None
    assert job["result"]["slide_count"] >= 1
    assert job["result"]["style_guide_id"] == created["deck_id"]

    deck = client.get(f"/v1/decks/{created['deck_id']}").json()
    assert deck["status"] == "completed"
    assert deck["slide_count"] == job["result"]["slide_count"]
    assert deck["template_count"] == job["result"]["template_count"]
    assert deck["style_guide_id"] == created["deck_id"]
    assert deck["error"] is None
    assert deck["job"]["status"] == "completed"

    guide_path = config.STYLE_GUIDES_DIR / f"{created['deck_id']}.json"
    templates_path = config.STYLE_GUIDES_DIR / f"{created['deck_id']}_templates.json"
    assert guide_path.is_file() and templates_path.is_file()
    guide = StyleGuide.model_validate_json(guide_path.read_text(encoding="utf-8"))
    assert guide.source_slide_count == deck["slide_count"]
    templates = load_templates(templates_path)
    assert len(templates) == deck["slide_count"]


def test_upload_rejects_non_pptx(client):
    response = client.post(
        "/v1/decks", files={"file": ("notes.txt", b"hello", "text/plain")}
    )
    assert response.status_code == 400
    assert "pptx" in response.json()["detail"]


def test_upload_rejects_legacy_ppt(client):
    response = client.post(
        "/v1/decks",
        files={"file": ("legacy.ppt", b"old binary", "application/octet-stream")},
    )
    assert response.status_code == 400
    detail = response.json()["detail"]
    assert "pptx" in detail
    assert "Save As" in detail


def _upload_bytes(client, name: str, data: bytes, mime: str):
    return client.post("/v1/decks", files={"file": (name, data, mime)})


def _tiny_pdf_bytes(tmp_path) -> bytes:
    import pymupdf

    path = tmp_path / "built.pdf"
    doc = pymupdf.open()
    for _ in range(2):
        page = doc.new_page(width=960, height=540)
        page.insert_text((72, 100), "Quarterly Review", fontsize=24, fontname="hebo")
    doc.save(str(path))
    doc.close()
    return path.read_bytes()


def test_upload_accepts_pdf_and_extracts_it(client, tmp_path):
    response = _upload_bytes(
        client, "report.pdf", _tiny_pdf_bytes(tmp_path), "application/pdf"
    )
    assert response.status_code == 202, response.text
    created = response.json()

    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job
    assert job["result"]["slide_count"] == 2

    deck = client.get(f"/v1/decks/{created['deck_id']}").json()
    assert deck["status"] == "completed"
    assert deck["source_name"] == "report.pdf"

    guide = StyleGuide.model_validate_json(
        (config.STYLE_GUIDES_DIR / f"{created['deck_id']}.json").read_text(
            encoding="utf-8"
        )
    )
    assert guide.source_file_type == ".pdf"


def test_upload_accepts_png_and_extracts_it(client, tmp_path):
    import pymupdf

    png_path = tmp_path / "hero.png"
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 320, 180))
    pix.set_rect(pix.irect, (0x0E, 0x0E, 0x10))
    pix.save(str(png_path))

    response = _upload_bytes(
        client, "hero.png", png_path.read_bytes(), "image/png"
    )
    assert response.status_code == 202, response.text
    created = response.json()

    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job
    assert job["result"]["slide_count"] == 1

    deck = client.get(f"/v1/decks/{created['deck_id']}").json()
    assert deck["status"] == "completed"

    guide = StyleGuide.model_validate_json(
        (config.STYLE_GUIDES_DIR / f"{created['deck_id']}.json").read_text(
            encoding="utf-8"
        )
    )
    assert guide.source_file_type == ".png"
    assert all(entry.provenance == "measured" for entry in guide.palette)


def test_upload_rejects_empty_file(client):
    response = client.post(
        "/v1/decks",
        files={"file": ("empty.pptx", b"", "application/octet-stream")},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_upload_rejects_oversized_file(client, fixture_deck, monkeypatch):
    monkeypatch.setattr(config, "MAX_UPLOAD_BYTES", 1000)
    response = _upload_deck(client, fixture_deck)
    assert response.status_code == 413
    assert list(config.UPLOADS_DIR.glob("deck_*")) == []


def test_upload_sanitizes_file_name(client, fixture_deck):
    response = _upload_deck(client, fixture_deck, name="../../quarterly review!.pptx")
    assert response.status_code == 202, response.text
    created = response.json()
    deck = client.get(f"/v1/decks/{created['deck_id']}").json()
    assert "/" not in deck["source_name"] and "\\" not in deck["source_name"]
    assert ".." not in deck["source_name"]
    assert deck["source_name"].endswith(".pptx")
    stored = list((config.UPLOADS_DIR / created["deck_id"]).iterdir())
    assert len(stored) == 1 and stored[0].name == deck["source_name"]
    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job


def test_invalid_pptx_fails_the_job_and_the_deck(client):
    response = client.post(
        "/v1/decks",
        files={"file": ("broken.pptx", b"not a zip archive", "application/octet-stream")},
    )
    assert response.status_code == 202, response.text
    created = response.json()
    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "failed"
    assert job["error"]
    assert job["result"] is None

    deck = client.get(f"/v1/decks/{created['deck_id']}").json()
    assert deck["status"] == "failed"
    assert deck["error"]


def test_unknown_deck_and_job_return_404(client):
    deck_response = client.get("/v1/decks/deck_does_not_exist")
    assert deck_response.status_code == 404
    job_response = client.get("/v1/jobs/job_does_not_exist")
    assert job_response.status_code == 404
    assert "job_does_not_exist" in job_response.json()["detail"]


def test_restart_marks_interrupted_jobs_and_decks_failed(client):
    db.insert_job("job_stale", "extraction")
    db.insert_deck("deck_stale", "x.pptx", config.UPLOADS_DIR / "x.pptx", "job_stale")
    db.update_job("job_stale", status="running", stage="parsing", progress=10)
    db.update_deck_status("deck_stale", "running")

    create_app()  # a fresh process: recovery runs at startup

    job = db.get_job("job_stale")
    assert job["status"] == "failed"
    assert "restarted" in job["error"]
    assert db.get_deck("deck_stale")["status"] == "failed"


# -- style guide ------------------------------------------------------------


def _prepare_style_guide(client, fixture_deck):
    """Upload the fixture deck and wait until its style guide is learned."""
    response = _upload_deck(client, fixture_deck)
    assert response.status_code == 202, response.text
    created = response.json()
    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job
    return created["deck_id"]


def test_style_guide_read_returns_learned_guide(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    response = client.get(f"/v1/style-guides/{deck_id}")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["overrides"] == []

    guide = StyleGuide.model_validate(body["style_guide"])
    deck = client.get(f"/v1/decks/{deck_id}").json()
    assert guide.source_slide_count == deck["slide_count"]
    assert guide.typography.title is not None
    assert guide.palette


def test_style_guide_patch_applies_persists_and_records(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    before = client.get(f"/v1/style-guides/{deck_id}").json()["style_guide"]
    accents_before = [e["hex"] for e in before["palette"] if e["usage"] == "accent"]
    assert accents_before  # the fixture deck yields accent entries

    patch = {
        "typography": {
            "title": {
                "font_family": "Georgia",
                "font_size_pt": 40,
                "color_hex": "#a1b2c3",
            }
        },
        "content_rules": {"max_bullets_per_slide": 4, "preferred_density": "low"},
        "layout_grid": {"margins_px": {"left_px": 96}},
        "palette": [{"hex": "#112233", "usage": "accent"}],
    }
    response = client.patch(f"/v1/style-guides/{deck_id}", json=patch)
    assert response.status_code == 200, response.text
    body = response.json()
    guide = body["style_guide"]

    assert guide["typography"]["title"]["font_family"] == "Georgia"
    assert guide["typography"]["title"]["font_size_pt"] == 40
    assert guide["typography"]["title"]["color_hex"] == "#A1B2C3"  # normalized
    # learned statistics survive an override untouched
    assert (
        guide["typography"]["title"]["provenance"]
        == before["typography"]["title"]["provenance"]
    )
    assert (
        guide["typography"]["title"]["sample_count"]
        == before["typography"]["title"]["sample_count"]
    )
    assert guide["content_rules"]["max_bullets_per_slide"] == 4
    assert guide["content_rules"]["preferred_density"] == "low"
    assert guide["layout_grid"]["margins_px"]["left_px"] == 96
    assert (
        guide["layout_grid"]["margins_px"]["top_px"]
        == before["layout_grid"]["margins_px"]["top_px"]
    )
    # every entry of the usage moves together, or the renderer could
    # silently keep showing the old color
    accent_after = [e["hex"] for e in guide["palette"] if e["usage"] == "accent"]
    assert accent_after == ["#112233"] * len(accents_before)

    assert body["applied"]["typography"]["title"]["color_hex"] == "#A1B2C3"
    assert len(body["overrides"]) == 1
    assert body["overrides"][0]["patch"]["content_rules"]["max_bullets_per_slide"] == 4

    # the effective guide is on disk, not only in the response
    on_disk = json.loads(
        (config.STYLE_GUIDES_DIR / f"{deck_id}.json").read_text(encoding="utf-8")
    )
    assert on_disk["content_rules"]["max_bullets_per_slide"] == 4

    # a second, partial patch keeps earlier overrides and appends to the trail
    second = client.patch(
        f"/v1/style-guides/{deck_id}",
        json={"content_rules": {"max_words_per_bullet": 10}},
    )
    assert second.status_code == 200, second.text
    second_body = second.json()
    assert second_body["style_guide"]["content_rules"]["max_bullets_per_slide"] == 4
    assert second_body["style_guide"]["content_rules"]["max_words_per_bullet"] == 10
    assert len(second_body["overrides"]) == 2

    re_get = client.get(f"/v1/style-guides/{deck_id}").json()
    assert re_get["style_guide"]["content_rules"]["max_words_per_bullet"] == 10
    assert len(re_get["overrides"]) == 2


def test_style_guide_patch_rejects_invalid_values(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)

    empty = client.patch(f"/v1/style-guides/{deck_id}", json={})
    assert empty.status_code == 400, empty.text

    bad_patches = [
        {"typography": {"title": {"color_hex": "#12345"}}},
        {"typography": {"title": {"font_weight": "heavy"}}},
        {"typography": {"title": {"font_size_pt": 0}}},
        {"content_rules": {"max_bullets_per_slide": 99}},
        {"content_rules": {"preferred_density": "extreme"}},
        {"layout_grid": {"margins_px": {"left_px": -1}}},
        {"unknown_field": True},
    ]
    for patch in bad_patches:
        response = client.patch(f"/v1/style-guides/{deck_id}", json=patch)
        assert response.status_code == 422, (patch, response.text)

    guide = client.get(f"/v1/style-guides/{deck_id}").json()
    assert guide["overrides"] == []  # rejected edits are never recorded


def test_style_guide_unknown_and_unsafe_ids_return_404(client):
    for style_guide_id in ("deck_missing", "..", "evil..name", "..%2Fdeckdna"):
        path = f"/v1/style-guides/{style_guide_id}"
        assert client.get(path).status_code == 404, style_guide_id
        response = client.patch(
            path, json={"content_rules": {"max_bullets_per_slide": 3}}
        )
        assert response.status_code == 404, style_guide_id


# -- outline ----------------------------------------------------------------


OUTLINE_NOTES = [
    "Launch the analytics workflow for enterprise teams",
    "42%: pilot conversion in the first month",
    "> Ship monthly, learn weekly - Platform team",
    "Improve activation through guided onboarding",
]


def test_outline_plans_and_persists(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    request = {
        "topic": "Q4 Launch Plan",
        "audience": "Leadership",
        "slide_count": 6,
        "style_guide_id": deck_id,
        "notes": OUTLINE_NOTES,
    }
    response = client.post("/v1/generations/outline", json=request)
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["outline_id"].startswith("outline_")
    assert body["style_guide_id"] == deck_id
    assert body["deck_title"] == "Q4 Launch Plan"
    assert body["audience"] == "Leadership"
    slides = body["slides"]
    assert body["slide_count"] == len(slides) == 6
    assert body["warnings"] == []
    assert slides[0]["slide_type"] == "title"
    assert slides[0]["title"] == "Q4 Launch Plan"
    assert [s["slide_number"] for s in slides] == list(range(1, 7))
    assert all(s["intent"] for s in slides)
    template_ids = {
        t.template_id
        for t in load_templates(config.STYLE_GUIDES_DIR / f"{deck_id}_templates.json")
    }
    assert {s["template_id"] for s in slides} <= template_ids

    # the plan is persisted for the generation step that follows approval
    stored = db.get_outline(body["outline_id"])
    assert stored is not None
    assert stored["style_guide_id"] == deck_id
    assert stored["deck_id"] == deck_id
    assert stored["brief"]["topic"] == "Q4 Launch Plan"
    assert stored["brief"]["notes"] == OUTLINE_NOTES
    assert len(stored["plan"]["slides"]) == 6

    # planning is deterministic: the same brief plans the same storyline
    again = client.post("/v1/generations/outline", json=request).json()
    assert again["outline_id"] != body["outline_id"]
    assert again["slides"] == slides


def test_outline_rejects_invalid_requests(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    base = {"topic": "Roadmap", "style_guide_id": deck_id, "slide_count": 5}
    cases = [
        {**base, "topic": ""},
        {**base, "topic": "   "},
        {**base, "topic": "x" * 201},
        {**base, "slide_count": 0},
        {**base, "slide_count": 51},
        {k: v for k, v in base.items() if k != "style_guide_id"},
        {**base, "surprise": 1},
        {**base, "notes": ["x" * 501]},
    ]
    for case in cases:
        response = client.post("/v1/generations/outline", json=case)
        assert response.status_code == 422, (case, response.text)

    unknown = client.post(
        "/v1/generations/outline", json={**base, "style_guide_id": "deck_nope"}
    )
    assert unknown.status_code == 404
    assert "deck_nope" in unknown.json()["detail"]


# -- generation -------------------------------------------------------------


def _plan_outline(client, deck_id, **overrides):
    request = {
        "topic": "Q4 Launch Plan",
        "audience": "Leadership",
        "slide_count": 6,
        "style_guide_id": deck_id,
        "notes": OUTLINE_NOTES,
    }
    request.update(overrides)
    response = client.post("/v1/generations/outline", json=request)
    assert response.status_code == 200, response.text
    return response.json()


def _generate(client, deck_id, outline_id, **overrides):
    request = {
        "style_guide_id": deck_id,
        "approved_outline_id": outline_id,
        "output_format": "pptx",
        "enable_critique": True,
    }
    request.update(overrides)
    return client.post("/v1/generations", json=request)


def _download(client, generation_id, artifact):
    return client.get(f"/v1/generations/{generation_id}/download/{artifact}")


def test_generation_runs_and_serves_downloads(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    outline = _plan_outline(client, deck_id)

    response = _generate(client, deck_id, outline["outline_id"])
    assert response.status_code == 202, response.text
    created = response.json()
    assert created["generation_id"].startswith("gen_")
    assert created["status"] == "queued"
    assert created["job_id"].startswith("job_")

    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job
    assert job["progress"] == 100
    result = job["result"]
    assert result["generation_id"] == created["generation_id"]
    assert result["outline_id"] == outline["outline_id"]
    assert result["style_guide_id"] == deck_id
    assert result["output_format"] == "pptx"
    assert result["deck_title"] == outline["deck_title"]
    assert result["slide_count"] == len(outline["slides"])
    assert result["generator"] == "deterministic"
    assert result["critique"]["requested"] is True
    assert result["critique"]["iterations"] >= 1
    assert result["critique"]["threshold"] > 0
    assert set(result["artifacts"]) == {"pptx", "html", "json", "critique"}
    assert (
        job["output_url"]
        == f"/v1/generations/{created['generation_id']}/download/pptx"
    )

    # the generated deck honors the approved outline exactly
    content = _download(client, created["generation_id"], "json")
    assert content.status_code == 200, content.text
    deck = GeneratedDeck.model_validate_json(content.content)
    assert [slide.model_dump() for slide in deck.plan.slides] == outline["slides"]
    assert deck.style_guide_id == deck_id

    # the preview is self-contained HTML of the same deck
    preview = _download(client, created["generation_id"], "html")
    assert preview.status_code == 200
    assert outline["deck_title"] in preview.text
    for slide in deck.slides:
        assert f'id="{slide.slide_id}"' in preview.text

    # the critique artifact is the full evidence-bearing report
    critique_response = _download(client, created["generation_id"], "critique")
    assert critique_response.status_code == 200
    report = CritiqueReport.model_validate_json(critique_response.content)
    assert report.score_before == result["critique"]["score_before"]
    assert report.score_after == result["critique"]["score_after"]
    assert report.stop_reason == result["critique"]["stop_reason"]

    # the pptx download is a real, openable PowerPoint file
    pptx = _download(client, created["generation_id"], "pptx")
    assert pptx.status_code == 200
    assert pptx.content[:2] == b"PK"
    presentation = Presentation(BytesIO(pptx.content))
    assert len(presentation.slides) == len(outline["slides"])
    assert "attachment" in pptx.headers["content-disposition"]


def test_generation_without_critique_skips_the_loop(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    outline = _plan_outline(client, deck_id)

    response = _generate(
        client, deck_id, outline["outline_id"], enable_critique=False
    )
    assert response.status_code == 202, response.text
    created = response.json()

    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job
    assert job["result"]["critique"] == {"requested": False}
    assert set(job["result"]["artifacts"]) == {"pptx", "html", "json"}
    assert _download(client, created["generation_id"], "critique").status_code == 404
    assert _download(client, created["generation_id"], "pptx").status_code == 200


def test_generation_rejects_unknown_outline_and_mismatched_guide(
    client, fixture_deck
):
    deck_id = _prepare_style_guide(client, fixture_deck)
    outline = _plan_outline(client, deck_id)

    unknown = _generate(client, deck_id, "outline_missing")
    assert unknown.status_code == 404
    assert "outline_missing" in unknown.json()["detail"]

    mismatch = _generate(client, "deck_other", outline["outline_id"])
    assert mismatch.status_code == 400
    assert "deck_other" in mismatch.json()["detail"]


def test_generation_rejects_invalid_requests(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    outline = _plan_outline(client, deck_id)
    base = {
        "style_guide_id": deck_id,
        "approved_outline_id": outline["outline_id"],
    }
    cases = [
        {k: v for k, v in base.items() if k != "style_guide_id"},
        {k: v for k, v in base.items() if k != "approved_outline_id"},
        {**base, "output_format": "pdf"},
        # pydantic's lax mode maps "true"/"false" strings, but arbitrary
        # text must not silently become a boolean
        {**base, "enable_critique": "maybe"},
        {**base, "surprise": 1},
    ]
    for case in cases:
        response = client.post("/v1/generations", json=case)
        assert response.status_code == 422, (case, response.text)


def test_generation_download_unknown_and_unsafe_ids(client):
    assert _download(client, "gen_missing", "pptx").status_code == 404
    assert _download(client, "gen..name", "pptx").status_code == 404
    assert (
        client.get(
            "/v1/generations/gen_x/download/..%2F..%2Fstyle_guides%2Fdeck.json"
        ).status_code
        == 404
    )
    assert _download(client, "gen_x", "zip").status_code == 404


# -- browser pages (Milestone 7 UI) -----------------------------------------


def test_home_page_lists_recent_decks(client, fixture_deck):
    empty = client.get("/")
    assert empty.status_code == 200
    assert "No decks uploaded yet" in empty.text

    deck_id = _prepare_style_guide(client, fixture_deck)
    home = client.get("/")
    assert home.status_code == 200, home.text
    assert "sample_deck.pptx" in home.text
    assert f"/style-guides/{deck_id}" in home.text
    assert "completed" in home.text


def test_deck_progress_page_shows_state_and_links(client, fixture_deck):
    response = _upload_deck(client, fixture_deck)
    assert response.status_code == 202, response.text
    created = response.json()

    page = client.get(f"/decks/{created['deck_id']}")
    assert page.status_code == 200, page.text
    assert "sample_deck.pptx" in page.text

    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job
    done = client.get(f"/decks/{created['deck_id']}")
    assert done.status_code == 200
    assert f"/style-guides/{created['deck_id']}" in done.text
    assert "reusable layouts" in done.text

    missing = client.get("/decks/deck_missing")
    assert missing.status_code == 404
    # quotes are HTML-escaped in the rendered page
    assert "No deck" in missing.text and "deck_missing" in missing.text
    assert client.get("/decks/gen..name").status_code == 404


def test_style_guide_page_shows_learned_values_and_edit_trail(
    client, fixture_deck
):
    deck_id = _prepare_style_guide(client, fixture_deck)
    page = client.get(f"/style-guides/{deck_id}")
    assert page.status_code == 200, page.text

    guide = client.get(f"/v1/style-guides/{deck_id}").json()["style_guide"]
    first_palette = guide["palette"][0]
    assert first_palette["hex"] in page.text
    assert first_palette["usage"] in page.text
    assert "No edits yet" in page.text

    records = load_templates(
        config.STYLE_GUIDES_DIR / f"{deck_id}_templates.json"
    )
    assert records[0].template_id in page.text
    assert "Generate a new deck in this style" in page.text

    patch = client.patch(
        f"/v1/style-guides/{deck_id}",
        json={"content_rules": {"max_bullets_per_slide": 4}},
    )
    assert patch.status_code == 200, patch.text
    edited = client.get(f"/style-guides/{deck_id}")
    assert "Change 1" in edited.text
    assert "max_bullets_per_slide" in edited.text

    assert client.get("/style-guides/deck_missing").status_code == 404
    assert client.get("/style-guides/gen..name").status_code == 404


def test_outline_page_lists_slides_and_approval_form(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    outline = _plan_outline(client, deck_id)

    page = client.get(f"/outlines/{outline['outline_id']}")
    assert page.status_code == 200, page.text
    assert outline["deck_title"] in page.text
    assert outline["outline_id"] in page.text
    assert f"/style-guides/{deck_id}" in page.text
    assert "Approve and generate" in page.text
    for slide in outline["slides"]:
        assert slide["template_id"] in page.text
        assert slide["intent"] in page.text

    assert client.get("/outlines/outline_missing").status_code == 404
    assert client.get("/outlines/evil..name").status_code == 404


def test_generation_page_renders_result_and_inline_preview(
    client, fixture_deck
):
    deck_id = _prepare_style_guide(client, fixture_deck)
    outline = _plan_outline(client, deck_id)
    created = _generate(client, deck_id, outline["outline_id"]).json()

    # whatever state the job is in, the page must render
    early = client.get(
        f"/generations/{created['generation_id']}",
        params={"job": created["job_id"]},
    )
    assert early.status_code == 200, early.text

    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job
    page = client.get(
        f"/generations/{created['generation_id']}",
        params={"job": created["job_id"]},
    )
    assert page.status_code == 200, page.text
    assert job["result"]["deck_title"] in page.text
    assert (
        f"/v1/generations/{created['generation_id']}/download/pptx" in page.text
    )
    assert "download/html?inline=1" in page.text
    assert "Quality check" in page.text

    # downloads keep the attachment disposition; the iframe needs inline
    attachment = _download(client, created["generation_id"], "html")
    assert "attachment" in attachment.headers["content-disposition"]
    inline = client.get(
        f"/v1/generations/{created['generation_id']}/download/html?inline=1"
    )
    assert "inline" in inline.headers["content-disposition"]
    assert inline.text == attachment.text

    # a page without a job id still gives a usable notice
    notice = client.get(f"/generations/{created['generation_id']}")
    assert notice.status_code == 200

    unknown_job = client.get(
        f"/generations/{created['generation_id']}",
        params={"job": "job_missing"},
    )
    assert unknown_job.status_code == 404
    assert client.get("/generations/gen..name").status_code == 404


def test_generation_page_shows_failed_job_error(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    outline = _plan_outline(client, deck_id)
    created = _generate(client, deck_id, outline["outline_id"]).json()
    job = _wait_for_job(client, created["job_id"])
    assert job["status"] == "completed", job

    # simulate a failure of an earlier run of the same-ish shape: swap the
    # stored job for a failed one and confirm the page reports it honestly
    db.update_job(created["job_id"], status="failed", error="disk full")
    page = client.get(
        f"/generations/{created['generation_id']}",
        params={"job": created["job_id"]},
    )
    assert page.status_code == 200
    assert "disk full" in page.text


def test_generation_job_reports_critique_iteration_stages(client, fixture_deck):
    deck_id = _prepare_style_guide(client, fixture_deck)
    outline = _plan_outline(client, deck_id)

    stored = db.get_outline(outline["outline_id"])
    style_guide = StyleGuide.model_validate_json(
        (config.STYLE_GUIDES_DIR / f"{deck_id}.json").read_text(encoding="utf-8")
    )
    templates = load_templates(config.STYLE_GUIDES_DIR / f"{deck_id}_templates.json")

    stages = []

    def report(stage, progress):
        stages.append((stage, progress))

    result = run_generation(
        "gen_stagecheck",
        outline["outline_id"],
        DeckBrief.model_validate(stored["brief"]),
        GenerationPlan.model_validate(stored["plan"]),
        style_guide,
        templates,
        True,
        report,
    )

    assert stages[0][0] == "generating_content"
    assert "critic_iteration_1" in [stage for stage, _ in stages]
    assert stages[-1][0] == "finalizing"
    assert [progress for _, progress in stages] == sorted(
        progress for _, progress in stages
    )
    assert result["generation_id"] == "gen_stagecheck"
    assert (config.GENERATED_DIR / "gen_stagecheck" / "deck.pptx").is_file()
    assert (config.GENERATED_DIR / "gen_stagecheck" / "preview.html").is_file()
    assert (config.GENERATED_DIR / "gen_stagecheck" / "content.json").is_file()
    assert (config.GENERATED_DIR / "gen_stagecheck" / "critique.json").is_file()
