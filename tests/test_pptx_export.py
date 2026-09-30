"""Offline tests for the editable PPTX export (blueprint §8.4, Milestone 5).
The artifact under test is the .pptx file itself, reopened with python-pptx:
native editable runs carrying the learned fonts, sizes and colors; slot
geometry converted to EMU on the learned canvas; shape names that carry the
HTML element ids; real bullet characters; dashed media placeholders instead
of fabricated charts; and the same z-order the HTML preview paints in. A
Windows-only acceptance test renders the exported file back to PNGs with
desktop PowerPoint, proving the file opens and lays out.
"""

from __future__ import annotations

import sys

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from core.generation import (
    BulletBlock,
    DeckBrief,
    GeneratedDeck,
    GenerationPlan,
    OutlineSlide,
    SlideContent,
    StatBlock,
    generate_deck,
)
from core.pptx_export import EMU_PER_PX, export_deck_pptx
from core.renderer import PowerPointComRenderer
from core.style_guide import (
    ContentRules,
    LayoutGrid,
    PaletteEntry,
    StyleGuide,
    Typography,
    TypographySpec,
)
from core.templates import (
    ROLE_BODY,
    ROLE_LABEL,
    ROLE_PICTURE,
    TemplateRecord,
    TemplateSlot,
)

EMU_PER_PT = 12700


# -- builders (mirrors tests/test_html_renderer.py) --------------------------


def _slot(role, x=0.05, y=0.1, w=0.9, h=0.3, size=18.0):
    return TemplateSlot(
        role=role, shape_type="TEXT_BOX", x=x, y=y, width=w, height=h,
        font_size_pt=size,
    )


def _template(slide_type, template_id, slots):
    return TemplateRecord(
        template_id=template_id,
        slide_type=slide_type,
        source_deck="synthetic",
        source_slide_number=1,
        aspect_ratio="16:9",
        slots=slots,
    )


def _library():
    return [
        _template("title", "tpl_title", [_slot("title", h=0.25)]),
        _template(
            "bullets",
            "tpl_bullets",
            [_slot("title", y=0.05, h=0.15), _slot(ROLE_BODY, y=0.3, h=0.5)],
        ),
        _template("stat", "tpl_stat", [_slot("stat", h=0.4)]),
        _template("quote", "tpl_quote", [_slot("quote", h=0.4)]),
        _template("section_divider", "tpl_divider", [_slot("title", h=0.3)]),
    ]


def _chart_template():
    return _template(
        "chart",
        "tpl_chart",
        [
            _slot("stat", x=0.0625, y=0.2778, w=0.3125, h=0.2037, size=96.0),
            _slot(ROLE_LABEL, x=0.0646, y=0.5, w=0.3125, h=0.0741, size=18.0),
        ],
    )


def _style_guide(palette=None, typo=None):
    return StyleGuide(
        deck_id="synthetic-deck",
        deck_name="synthetic",
        extracted_at="2026-01-01T00:00:00+00:00",
        source_file_type="pptx",
        source_slide_count=6,
        aspect_ratio="16:9",
        palette=palette or [],
        typography=typo or Typography(),
        layout_grid=LayoutGrid(slide_width_px=1920.0, slide_height_px=1080.0),
        content_rules=ContentRules(),
    )


NOTES = [
    "Revenue growth strategy",
    "Retention plans for next year",
    "Growth doubled after we changed the pricing model",
    "Support load fell with self serve docs",
    "> We ship every week \u2014 the engineering team",
    "82% \u2014 of teams ship weekly",
]


def _brief(notes=NOTES, slide_count=5):
    return DeckBrief(
        topic="Quarterly Review",
        audience="Leadership",
        goal="Align on next quarter",
        slide_count=slide_count,
        notes=list(notes),
    )


def _deck_from(slides):
    plan = GenerationPlan(
        deck_title="Synthetic deck",
        slides=[
            OutlineSlide(
                slide_number=slide.slide_number,
                intent="test slide",
                slide_type=slide.slide_type,
                template_id=slide.template_id,
                title=slide.title,
            )
            for slide in slides
        ],
    )
    return GeneratedDeck(brief=DeckBrief(topic="Synthetic deck"), plan=plan, slides=slides)


def _deterministic_deck():
    return generate_deck(_brief(), _library(), _style_guide())


# -- helpers -----------------------------------------------------------------


def _export_open(deck, guide, templates, tmp_path, name="deck.pptx"):
    path = export_deck_pptx(deck, guide, templates, tmp_path / name)
    assert path.is_file()
    return Presentation(str(path))


def _shape(prs, slide_index, name):
    for shape in prs.slides[slide_index].shapes:
        if shape.name == name:
            return shape
    names = [shape.name for shape in prs.slides[slide_index].shapes]
    raise AssertionError(f"{name!r} not found; slide has {names}")


def _pPr(paragraph):
    p_pr = paragraph._p.find(qn("a:pPr"))
    assert p_pr is not None
    return p_pr


# -- file under test ---------------------------------------------------------


class TestExportDeckPptx:
    def test_slide_canvas_matches_the_learned_grid(self, tmp_path):
        prs = _export_open(
            _deterministic_deck(), _style_guide(), _library(), tmp_path
        )
        assert prs.slide_width == round(1920.0 * EMU_PER_PX)
        assert prs.slide_height == round(1080.0 * EMU_PER_PX)

    def test_every_slide_is_exported_with_the_palette_background(self, tmp_path):
        guide = _style_guide(
            palette=[PaletteEntry(hex="#101820", usage="background", frequency=5)]
        )
        prs = _export_open(_deterministic_deck(), guide, _library(), tmp_path)
        assert len(prs.slides) == 5
        for slide in prs.slides:
            assert slide.background.fill.fore_color.rgb == RGBColor.from_string("101820")

    def test_element_names_are_stable(self, tmp_path):
        prs = _export_open(
            _deterministic_deck(), _style_guide(), _library(), tmp_path
        )
        names = [shape.name for slide in prs.slides for shape in slide.shapes]
        assert len(names) == len(set(names))
        assert {
            "slide_01-title",
            "slide_01-subtitle",
            "slide_02-quote",
            "slide_03-stat",
            "slide_04-body",
            "slide_05-title",
        } <= set(names)

    def test_title_is_a_native_editable_run(self, tmp_path):
        typo = Typography(
            title=TypographySpec(
                font_family="Trebuchet MS", font_size_pt=44.0, color_hex="#FFD166",
                sample_count=3,
            ),
            body=TypographySpec(font_family="Georgia", sample_count=3),
        )
        guide = _style_guide(typo=typo)
        template = _template(
            "title", "tpl_title", [_slot("title", h=0.25, size=44.0)]
        )
        slide = SlideContent(
            slide_id="slide_01", slide_number=1, slide_type="title",
            template_id="tpl_title", title="Quarterly Review",
        )
        prs = _export_open(_deck_from([slide]), guide, [template], tmp_path)
        shape = _shape(prs, 0, "slide_01-title")
        assert shape.has_text_frame
        assert shape.text_frame.text == "Quarterly Review"
        assert shape.text_frame.word_wrap is True
        assert shape.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
        run = shape.text_frame.paragraphs[0].runs[0]
        assert run.font.size == Pt(44)
        assert run.font.bold is True
        assert run.font.italic is False
        assert run.font.name == "Trebuchet MS"
        assert run.font.color.rgb == RGBColor.from_string("FFD166")

    def test_geometry_follows_template_fractions(self, tmp_path):
        prs = _export_open(
            _deterministic_deck(), _style_guide(), _library(), tmp_path
        )
        shape = _shape(prs, 3, "slide_04-body")
        # tpl_bullets body slot: x=0.05 y=0.3 w=0.9 h=0.5 on a 1920x1080 grid
        assert shape.left == round(0.05 * 1920 * EMU_PER_PX)
        assert shape.top == round(0.3 * 1080 * EMU_PER_PX)
        assert shape.width == round(0.9 * 1920 * EMU_PER_PX)
        assert shape.height == round(0.5 * 1080 * EMU_PER_PX)

    def test_bullets_are_native_bullet_paragraphs(self, tmp_path):
        guide = _style_guide(
            palette=[PaletteEntry(hex="#FFC857", usage="accent", frequency=2)]
        )
        prs = _export_open(_deterministic_deck(), guide, _library(), tmp_path)
        shape = _shape(prs, 3, "slide_04-body")
        paragraphs = shape.text_frame.paragraphs
        assert [p.text for p in paragraphs] == [
            "Growth doubled after we changed the pricing model",
            "Support load fell with self serve docs",
        ]
        marl = round(1.05 * 18.0 * EMU_PER_PT)
        for paragraph in paragraphs:
            p_pr = _pPr(paragraph)
            assert p_pr.get("marL") == str(marl)
            assert p_pr.get("indent") == str(-marl)
            bu_char = p_pr.find(qn("a:buChar"))
            assert bu_char is not None and bu_char.get("char") == "\u2022"
            accent = p_pr.find(qn("a:buClr") + "/" + qn("a:srgbClr"))
            assert accent is not None and accent.get("val") == "FFC857"
            assert paragraph.runs[0].font.size == Pt(18)

    def test_stat_context_lands_on_the_caption_slot(self, tmp_path):
        library = [_chart_template()]
        slide = SlideContent(
            slide_id="slide_01", slide_number=1, slide_type="chart",
            template_id="tpl_chart", title="Weekly shipping cadence",
            content=[StatBlock(value="82%", context="of teams ship weekly")],
        )
        prs = _export_open(_deck_from([slide]), _style_guide(), library, tmp_path)
        stat = _shape(prs, 0, "slide_01-stat")
        caption = _shape(prs, 0, "slide_01-stat-caption")
        assert stat.text_frame.text == "82%"
        assert stat.text_frame.paragraphs[0].runs[0].font.size == Pt(96)
        assert caption.text_frame.text == "of teams ship weekly"
        assert caption.top == round(0.5 * 1080 * EMU_PER_PX)

    def test_media_slots_render_dashed_placeholders_not_fabricated_charts(
        self, tmp_path
    ):
        template = _template(
            "bullets",
            "tpl_pic",
            [_slot("title", h=0.15), _slot(ROLE_PICTURE, y=0.3, h=0.6)],
        )
        slide = SlideContent(
            slide_id="slide_01", slide_number=1, slide_type="bullets",
            template_id="tpl_pic", title="Has a picture",
            content=[BulletBlock(text="One point")],
        )
        prs = _export_open(_deck_from([slide]), _style_guide(), [template], tmp_path)
        shape = _shape(prs, 0, "slide_01-picture")
        assert shape.fill.fore_color.rgb == RGBColor.from_string("F0F0F0")
        assert shape.line.color.rgb == RGBColor.from_string("B9B9B9")
        assert shape.line.dash_style == MSO_LINE_DASH_STYLE.DASH
        assert shape.text_frame.text == "PICTURE"
        assert shape.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
        assert shape.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
        for slide in prs.slides:
            for shape in slide.shapes:
                assert not shape.has_chart and not shape.has_table

    def test_card_shapes_paint_first_with_fill_and_learned_radius(self, tmp_path):
        guide = _style_guide()
        guide.element_treatments.card_corner_radius_px = 24.0
        template = _template(
            "comparison",
            "tpl_cmp",
            [
                _slot(ROLE_BODY, x=0.55, y=0.3, w=0.4, h=0.5),
                _slot("card", x=0.05, y=0.3, w=0.4, h=0.5),
            ],
        )
        template.slots[1].fill_hex = "#EEEEEE"
        slide = SlideContent(
            slide_id="slide_01", slide_number=1, slide_type="comparison",
            template_id="tpl_cmp",
            content=[BulletBlock(text="Option A is cheaper")],
        )
        prs = _export_open(_deck_from([slide]), guide, [template], tmp_path)
        slide_shapes = list(prs.slides[0].shapes)
        card = _shape(prs, 0, "slide_01-card")
        assert slide_shapes.index(card) == 0  # background paints before text
        assert card.fill.fore_color.rgb == RGBColor.from_string("EEEEEE")
        # adjustment is radius / min(card side) on the 1920x1080 canvas
        assert card.adjustments[0] == pytest.approx(24.0 / 540.0, abs=1e-3)

    def test_decorative_boxes_without_fill_are_skipped(self, tmp_path):
        template = _template(
            "bullets",
            "tpl_plain",
            [_slot("title", h=0.15), _slot("decorative", x=0.0, y=0.0, w=0.1, h=0.1)],
        )
        slide = SlideContent(
            slide_id="slide_01", slide_number=1, slide_type="bullets",
            template_id="tpl_plain", title="Plain",
        )
        prs = _export_open(_deck_from([slide]), _style_guide(), [template], tmp_path)
        names = [shape.name for shape in prs.slides[0].shapes]
        assert "slide_01-decorative" not in names

    def test_export_is_structurally_deterministic(self, tmp_path):
        first = _export_open(
            _deterministic_deck(), _style_guide(), _library(), tmp_path, "a.pptx"
        )
        second = _export_open(
            _deterministic_deck(), _style_guide(), _library(), tmp_path, "b.pptx"
        )
        for slide_a, slide_b in zip(first.slides, second.slides):
            fingerprint_a = [
                (s.name, s.left, s.top, s.width, s.height) for s in slide_a.shapes
            ]
            fingerprint_b = [
                (s.name, s.left, s.top, s.width, s.height) for s in slide_b.shapes
            ]
            assert fingerprint_a == fingerprint_b


# -- PowerPoint acceptance ---------------------------------------------------


def _png_size(path):
    from PIL import Image

    with Image.open(path) as image:
        return image.size


@pytest.mark.skipif(
    sys.platform != "win32" or not PowerPointComRenderer().available(),
    reason="desktop PowerPoint not available on this machine",
)
def test_powerpoint_opens_and_renders_the_exported_deck(tmp_path):
    path = export_deck_pptx(
        _deterministic_deck(), _style_guide(), _library(), tmp_path / "deck.pptx"
    )
    pages = PowerPointComRenderer().render_pptx(path, tmp_path / "slides")
    assert [page.name for page in pages] == [
        f"slide_{i:03d}.png" for i in range(1, 6)
    ]
    for page in pages:
        assert page.is_file()
        assert _png_size(page) == (1920, 1080)
