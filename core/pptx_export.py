"""Editable PPTX export (blueprint §8.4, Milestone 5): writes a GeneratedDeck
as a native PowerPoint file — real text boxes with editable runs, real shape
fills, dashed media placeholders.

Planning is shared with the HTML preview through core.layout so the two
outputs cannot drift: same slots, same fonts, same text parts. Geometry
converts canvas fractions to English Metric Units at 96 px per inch, which
makes the slide exactly the learned canvas size, so both outputs land on the
same pixel grid. Element ids become shape names (slide_02-title) so critique
fixes can target elements deterministically. Media slots that have no real
asset behind them render as labeled dashed placeholders — never fabricated
charts or tables.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from core.generation import GeneratedDeck
from core.layout import (
    DEFAULT_BACKGROUND,
    DEFAULT_TEXT,
    LINE_HEIGHT,
    RenderBox,
    canvas_size,
    font_family,
    palette_color,
    plan_slide,
)
from core.style_guide import StyleGuide
from core.templates import ROLE_CARD, TemplateRecord

EMU_PER_PX = 9525  # 96 px/inch, 914400 EMU/inch
EMU_PER_PT = 12700

# HTML placeholder styling (rgba over white), precomputed for PowerPoint.
_PLACEHOLDER_BORDER = "B9B9B9"
_PLACEHOLDER_FILL = "F0F0F0"
_PLACEHOLDER_TEXT = "707070"

_BULLET_CHAR = "\u2022"
_BULLET_INDENT_EM = 1.05  # matches ul.bullets li padding-left in the HTML CSS
_BULLET_SUCCESSORS = ("a:tabLst", "a:defRPr", "a:extLst")


def export_deck_pptx(
    deck: GeneratedDeck,
    style_guide: StyleGuide,
    templates: Sequence[TemplateRecord],
    output_path: Path,
) -> Path:
    canvas_w, canvas_h = canvas_size(style_guide)
    prs = Presentation()
    prs.slide_width = Emu(round(canvas_w * EMU_PER_PX))
    prs.slide_height = Emu(round(canvas_h * EMU_PER_PX))
    blank_layout = prs.slide_layouts[6]
    by_id = {record.template_id: record for record in templates}

    for slide in deck.slides:
        pptx_slide = prs.slides.add_slide(blank_layout)
        _paint_background(pptx_slide, style_guide)
        boxes, _ = plan_slide(slide, by_id.get(slide.template_id), style_guide)
        for box in boxes:
            if box.kind == "shape":
                _add_shape_box(pptx_slide, box, style_guide, canvas_w, canvas_h)
            elif box.kind == "placeholder":
                _add_placeholder_box(pptx_slide, box, canvas_w, canvas_h)
            else:
                _add_text_box(pptx_slide, box, style_guide, canvas_w, canvas_h)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def _paint_background(slide, style_guide: StyleGuide) -> None:
    background = (
        palette_color(style_guide, "background", DEFAULT_BACKGROUND)
        or DEFAULT_BACKGROUND
    )
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(background)


def _rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value.lstrip("#").upper())


def _geom(slot, canvas_w: float, canvas_h: float):
    return (
        Emu(round(slot.x * canvas_w * EMU_PER_PX)),
        Emu(round(slot.y * canvas_h * EMU_PER_PX)),
        Emu(round(slot.width * canvas_w * EMU_PER_PX)),
        Emu(round(slot.height * canvas_h * EMU_PER_PX)),
    )


def _add_shape_box(
    slide, box: RenderBox, style_guide: StyleGuide, canvas_w: float, canvas_h: float
) -> None:
    is_card = box.base == ROLE_CARD
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if is_card else MSO_SHAPE.RECTANGLE,
        *_geom(box.slot, canvas_w, canvas_h),
    )
    shape.name = box.element_id
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(box.fill_hex or DEFAULT_BACKGROUND)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if is_card:
        radius_px = style_guide.element_treatments.card_corner_radius_px
        if radius_px is not None and radius_px > 0:
            min_dim = min(
                box.slot.width * canvas_w, box.slot.height * canvas_h
            )
            if min_dim > 0:
                shape.adjustments[0] = min(radius_px / min_dim, 0.5)


def _add_placeholder_box(
    slide, box: RenderBox, canvas_w: float, canvas_h: float
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, *_geom(box.slot, canvas_w, canvas_h)
    )
    shape.name = box.element_id
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(_PLACEHOLDER_FILL)
    shape.line.color.rgb = _rgb(_PLACEHOLDER_BORDER)
    shape.line.width = Pt(1.5)  # the HTML dashed border is 2px
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shape.shadow.inherit = False

    text_frame = shape.text_frame
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = box.parts[0].text.upper()
    font = run.font
    if box.font_pt is not None:
        font.size = Pt(box.font_pt)
    font.bold = False
    font.italic = False
    font.color.rgb = _rgb(_PLACEHOLDER_TEXT)


def _add_text_box(
    slide, box: RenderBox, style_guide: StyleGuide, canvas_w: float, canvas_h: float
) -> None:
    shape = slide.shapes.add_textbox(*_geom(box.slot, canvas_w, canvas_h))
    shape.name = box.element_id
    text_frame = shape.text_frame
    text_frame.word_wrap = True
    if "vcenter" in box.css_classes:
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    family = font_family(style_guide, "title" if box.kind == "title" else "body")
    color = box.color or DEFAULT_TEXT
    accent = palette_color(style_guide, "accent", None)

    for index, part in enumerate(box.parts):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        run = paragraph.add_run()
        run.text = part.text
        font = run.font
        font_pt = part.font_pt if part.font_pt is not None else box.font_pt
        if font_pt is not None:
            font.size = Pt(font_pt)
        font.bold = _is_bold(box, part.part_kind)
        font.italic = box.slot.italic is True or part.part_kind == "quote-text"
        font.color.rgb = _rgb(color)
        if family:
            font.name = family

        if part.part_kind == "bullet":
            paragraph.line_spacing = LINE_HEIGHT["default"]
            _set_bullet(paragraph, accent, font_pt)
        elif part.part_kind == "stat-value":
            paragraph.line_spacing = LINE_HEIGHT["stat-value"]
        elif part.part_kind == "stat-label":
            paragraph.line_spacing = LINE_HEIGHT["default"]
            if part.font_pt is not None:
                paragraph.space_after = Pt(0.25 * part.font_pt)
        elif part.part_kind == "stat-context":
            paragraph.line_spacing = LINE_HEIGHT["default"]
            if part.font_pt is not None:
                paragraph.space_before = Pt(0.35 * part.font_pt)
        elif part.part_kind == "quote-attribution":
            paragraph.line_spacing = LINE_HEIGHT["default"]
            if part.font_pt is not None:
                paragraph.space_before = Pt(0.5 * part.font_pt)
        else:
            paragraph.line_spacing = (
                LINE_HEIGHT["title"] if box.kind == "title" else LINE_HEIGHT["default"]
            )


def _is_bold(box: RenderBox, part_kind: str) -> bool:
    if part_kind == "stat-value":
        return True  # .stat-value is font-weight 700 in the HTML CSS
    if box.kind == "title" and box.slot.bold is not False:
        return True  # h1.el is font-weight 600 in the HTML CSS
    return box.slot.bold is True


def _set_bullet(paragraph, accent_hex: Optional[str], font_pt: Optional[float]) -> None:
    """Give a paragraph a native bullet char; coordinates mirror the HTML
    list (bullet at the box edge, text indented 1.05em)."""
    p_pr = paragraph._p.get_or_add_pPr()
    if font_pt is not None:
        mar_l = round(_BULLET_INDENT_EM * font_pt * EMU_PER_PT)
        p_pr.set("marL", str(mar_l))
        p_pr.set("indent", str(-mar_l))
    if accent_hex:
        bu_clr = p_pr.makeelement(qn("a:buClr"), {})
        srgb = p_pr.makeelement(qn("a:srgbClr"), {"val": accent_hex.lstrip("#").upper()})
        bu_clr.append(srgb)
        p_pr.insert_element_before(
            bu_clr,
            "a:buSzTx", "a:buSzPct", "a:buSzPts",
            "a:buFontTx", "a:buFont",
            "a:buNone", "a:buAutoNum", "a:buChar",
            *_BULLET_SUCCESSORS,
        )
    bu_char = p_pr.makeelement(qn("a:buChar"), {"char": _BULLET_CHAR})
    p_pr.insert_element_before(bu_char, *_BULLET_SUCCESSORS)
